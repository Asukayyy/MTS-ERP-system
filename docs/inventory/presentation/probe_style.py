# -*- coding: utf-8 -*-
"""探测参考 PPT 矩形12/14 的填充、线条、字色字号。"""
from pptx import Presentation
from pptx.util import Emu, Pt
from lxml import etree

prs = Presentation(r"c:\Users\23681\Desktop\课程设计参考资料\计划管理模块.pptx")
s3 = prs.slides[2]
for sh in s3.shapes:
    if sh.name in ("矩形 12", "矩形 14"):
        print("====", sh.name)
        print(etree.tostring(sh._element, pretty_print=True).decode()[:3000])
