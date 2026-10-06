# -*- coding: utf-8 -*-
"""精简列出 1.pptx 每页形状与文本，定位可复用的源页。"""
from pptx import Presentation
from pptx.util import Emu

prs = Presentation(r"c:\Users\23681\Desktop\课程设计参考资料\1.pptx")
print("size", round(Emu(prs.slide_width).inches, 2), round(Emu(prs.slide_height).inches, 2))
for si, slide in enumerate(prs.slides, 1):
    print(f"--- slide {si} layout={slide.slide_layout.name!r} n={len(slide.shapes)}")
    for sh in slide.shapes:
        t = ""
        if sh.has_text_frame:
            t = " / ".join(p.text for p in sh.text_frame.paragraphs if p.text)[:70]
        pic = " [PIC]" if sh.shape_type == 13 else ""
        grp = " [GRP]" if sh.shape_type == 6 else ""
        print(f"    {sh.name!r}{pic}{grp} {t!r}")
