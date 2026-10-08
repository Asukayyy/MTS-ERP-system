# -*- coding: utf-8 -*-
"""提取 计划管理模块.pptx 里 slide1(封面)/slide2(目录)/slide10(截图) 的图片。"""
import os
from pptx import Presentation

prs = Presentation(r"c:\Users\23681\Desktop\课程设计参考资料\计划管理模块.pptx")
out = r"C:\Users\23681\AppData\Local\Temp\ref_imgs2"
os.makedirs(out, exist_ok=True)
for idx in (0, 1, 9):
    slide = prs.slides[idx]
    for sh in slide.shapes:
        if sh.shape_type == 13:
            img = sh.image
            fn = f"slide{idx+1}_{sh.shape_id}.{img.ext}"
            with open(os.path.join(out, fn), "wb") as f:
                f.write(img.blob)
            print(fn, img.size, sh.left, sh.top, sh.width, sh.height)
