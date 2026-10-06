# -*- coding: utf-8 -*-
"""检查参考 PPT 各页 组合1 内的可见文字（顶栏章节标识）。"""
from pptx import Presentation

prs = Presentation(r"c:\Users\23681\Desktop\课程设计参考资料\计划管理模块.pptx")
for idx in (0, 1, 2, 9, 11, 12):
    slide = prs.slides[idx]
    print(f"--- slide {idx+1}")
    for sh in slide.shapes:
        if sh.shape_type == 6:
            texts = []
            for sub in sh.shapes:
                if sub.has_text_frame and sub.text_frame.text.strip():
                    texts.append(sub.text_frame.text.strip()[:30])
            print(f"   group {sh.name!r}: {texts}")
