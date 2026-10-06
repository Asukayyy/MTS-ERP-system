# -*- coding: utf-8 -*-
"""回读校验 库存管理模块.pptx：页数、文本、图片、越界检查。"""
from pptx import Presentation
from pptx.util import Emu

PATH = r"c:\Users\23681\Desktop\课程设计参考资料\MTS-ERP-system-main\docs\inventory\presentation\库存管理模块.pptx"
prs = Presentation(PATH)
SW, SH = prs.slide_width, prs.slide_height
print(f"slides={len(prs.slides)} size={Emu(SW).inches:.2f}x{Emu(SH).inches:.2f}")
for i, slide in enumerate(prs.slides, 1):
    pics, texts = [], []
    for sh in slide.shapes:
        if sh.shape_type == 13:
            pics.append(sh.image.size)
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip().replace("\n", " / ")[:46])
        # 越界
        if sh.left is not None and sh.top is not None and sh.width and sh.height:
            if sh.left < -50000 or sh.top < -50000 or sh.left + sh.width > SW + 50000 or sh.top + sh.height > SH + 50000:
                if sh.shape_type == 13 or (sh.has_text_frame and sh.text_frame.text.strip()):
                    print(f"  !! slide{i} OUT-OF-PAGE {sh.name}")
    head = texts[0] if texts else "(no text)"
    print(f"[{i:02d}] pics={len(pics)} {pics[:2]} | {head}")
